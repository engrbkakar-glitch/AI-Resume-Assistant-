import json
import os
import re
from io import BytesIO

import streamlit as st
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document


st.set_page_config(
    page_title="Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
)

MAX_FILE_MB = 10
MODEL_NAME = "gemini-2.5-flash"


def get_api_key():
    """Read the Gemini API key from Streamlit secrets or the environment."""
    try:
        key = st.secrets.get("GEMINI_API_KEY")
        if key:
            return key
    except Exception:
        pass

    return os.getenv("GEMINI_API_KEY")


def extract_text(uploaded_file):
    """Extract text from PDF, DOCX, or TXT resumes."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if len(data) > MAX_FILE_MB * 1024 * 1024:
        raise ValueError(f"File is larger than {MAX_FILE_MB} MB.")

    if name.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = "\n".join(pages)

    elif name.endswith(".docx"):
        document = Document(BytesIO(data))
        parts = [p.text for p in document.paragraphs]

        # Also read simple table content because some resumes use tables.
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        text = "\n".join(parts)

    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="replace")

    else:
        raise ValueError("Unsupported file type. Please upload PDF, DOCX, or TXT.")

    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    if not text:
        raise ValueError(
            "No readable text was found. If your PDF is a scanned image, "
            "OCR is required before this app can analyze it."
        )

    return text


def analyze_resume(resume_text, job_description=""):
    """Ask Gemini to evaluate the resume and return structured JSON."""
    api_key = get_api_key()
    if not api_key:
        raise RuntimeError(
            "Gemini API key not found. Add GEMINI_API_KEY to Streamlit Secrets "
            "or set it as an environment variable."
        )

    client = genai.Client(api_key=api_key)

    jd_context = (
        job_description.strip()
        if job_description.strip()
        else "No job description was provided. Score general ATS readiness."
    )

    prompt = f"""
You are an expert ATS resume reviewer and career-document analyst.

Analyze the resume below. Give a practical ATS-readiness score from 0 to 100.
This is a heuristic score, not a score from a real ATS vendor.

If a job description is provided, evaluate keyword alignment against it.
Do NOT invent qualifications, experience, employers, degrees, dates, skills,
certifications, or achievements that are not present in the resume.

Scoring guidance:
- 25 points: ATS-safe structure and formatting
- 20 points: keyword and skills alignment
- 20 points: experience/achievement clarity
- 15 points: section completeness and organization
- 10 points: measurable impact and strong action language
- 10 points: readability, consistency, and common ATS issues

Return ONLY valid JSON matching this structure:
{{
  "ats_score": 0,
  "score_label": "Poor|Needs Improvement|Good|Very Good|Excellent",
  "summary": "Short overall assessment.",
  "formatting": {{
    "score": 0,
    "findings": ["finding 1", "finding 2"]
  }},
  "keywords": {{
    "score": 0,
    "matched_keywords": ["keyword"],
    "missing_or_weak_keywords": ["keyword"],
    "findings": ["finding"]
  }},
  "content": {{
    "score": 0,
    "findings": ["finding"]
  }},
  "section_issues": ["issue"],
  "priority_improvements": [
    {{
      "priority": "High|Medium|Low",
      "issue": "Specific issue",
      "recommendation": "Specific improvement",
      "example": "Optional example using only information already present in the resume."
    }}
  ],
  "ats_checklist": {{
    "standard_headings": true,
    "simple_layout": true,
    "contact_information_clear": true,
    "date_format_consistent": true,
    "action_verbs_used": true,
    "quantified_achievements": true,
    "keyword_alignment": true
  }}
}}

Rules:
- ats_score must be an integer from 0 to 100.
- Sub-scores must be integers from 0 to 100.
- Keep findings concise and actionable.
- If no job description is supplied, do not claim that a keyword is required by a
  specific employer or role.
- Treat tables, columns, icons, graphics, headers/footers, unusual symbols,
  and decorative formatting as potential ATS risks only when the extracted text
  gives evidence of them; otherwise say that text extraction cannot verify them.

JOB DESCRIPTION:
{jd_context}

RESUME:
{resume_text}
"""

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )

    raw = response.text.strip()

    # Defensive cleanup in case a model returns a fenced JSON block.
    raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.IGNORECASE)
    raw = re.sub(r"\s*```$", "", raw)

    try:
        result = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise RuntimeError("Gemini returned an invalid JSON response.") from exc

    validate_result(result)
    return result


def validate_result(result):
    """Basic validation so bad model output does not break the UI."""
    if not isinstance(result, dict):
        raise ValueError("Unexpected analysis format.")

    score = result.get("ats_score")
    if not isinstance(score, int) or not 0 <= score <= 100:
        raise ValueError("Gemini returned an invalid ATS score.")

    for key in ("formatting", "keywords", "content"):
        section = result.get(key, {})
        section_score = section.get("score")
        if not isinstance(section_score, int) or not 0 <= section_score <= 100:
            raise ValueError(f"Gemini returned an invalid {key} score.")


def score_color(score):
    if score >= 80:
        return "🟢"
    if score >= 60:
        return "🟡"
    return "🔴"


st.title("📄 Resume ATS Analyzer")
st.caption(
    "Upload your resume to get a heuristic ATS-readiness score and practical "
    "improvement suggestions powered by Gemini Flash."
)

with st.sidebar:
    st.header("How it works")
    st.markdown(
        "1. Upload a resume\n"
        "2. Optionally paste a job description\n"
        "3. Gemini analyzes ATS readiness\n"
        "4. Review the score and prioritized improvements"
    )
    st.info(
        "Privacy note: the resume text is sent to the Gemini API for analysis. "
        "Do not upload documents containing information you are not comfortable "
        "sending to the configured AI service."
    )

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx", "txt"],
    help="Maximum file size: 10 MB.",
)

job_description = st.text_area(
    "Job description (optional, but recommended)",
    height=220,
    placeholder="Paste the target job description here for keyword matching...",
)

if st.button("🔍 Analyze Resume", type="primary", use_container_width=True):
    if uploaded_file is None:
        st.warning("Please upload a resume first.")
    else:
        try:
            with st.spinner("Reading your resume and analyzing it with Gemini..."):
                resume_text = extract_text(uploaded_file)
                result = analyze_resume(resume_text, job_description)

            st.session_state["analysis"] = result

        except Exception as exc:
            st.error(f"Analysis failed: {exc}")

result = st.session_state.get("analysis")

if result:
    st.divider()

    col1, col2, col3, col4 = st.columns(4)
    col1.metric("ATS Score", f"{result['ats_score']}/100")
    col2.metric("Formatting", f"{result['formatting']['score']}/100")
    col3.metric("Keywords", f"{result['keywords']['score']}/100")
    col4.metric("Content", f"{result['content']['score']}/100")

    st.subheader(
        f"{score_color(result['ats_score'])} {result.get('score_label', 'ATS Assessment')}"
    )
    st.write(result.get("summary", ""))

    st.subheader("🎯 Priority Improvements")
    improvements = result.get("priority_improvements", [])
    if improvements:
        for item in improvements:
            priority = item.get("priority", "Medium")
            st.markdown(f"**{priority} — {item.get('issue', '')}**")
            st.write(item.get("recommendation", ""))
            if item.get("example"):
                st.caption(f"Example: {item['example']}")
    else:
        st.write("No priority improvements were returned.")

    st.subheader("🔑 Keyword Analysis")
    matched = result["keywords"].get("matched_keywords", [])
    missing = result["keywords"].get("missing_or_weak_keywords", [])

    c1, c2 = st.columns(2)
    with c1:
        st.markdown("**Matched / detected keywords**")
        st.write(", ".join(matched) if matched else "None detected.")
    with c2:
        st.markdown("**Missing / weak keywords**")
        st.write(", ".join(missing) if missing else "None identified.")

    st.write(result["keywords"].get("findings", []))

    st.subheader("🧾 Formatting Findings")
    for finding in result["formatting"].get("findings", []):
        st.markdown(f"- {finding}")

    st.subheader("✍️ Content Findings")
    for finding in result["content"].get("findings", []):
        st.markdown(f"- {finding}")

    st.subheader("📌 Section Issues")
    for issue in result.get("section_issues", []):
        st.markdown(f"- {issue}")

    st.subheader("✅ ATS Checklist")
    checklist = result.get("ats_checklist", {})
    checklist_items = {
        "Standard headings": checklist.get("standard_headings"),
        "Simple layout": checklist.get("simple_layout"),
        "Contact information clear": checklist.get("contact_information_clear"),
        "Consistent dates": checklist.get("date_format_consistent"),
        "Action verbs": checklist.get("action_verbs_used"),
        "Quantified achievements": checklist.get("quantified_achievements"),
        "Keyword alignment": checklist.get("keyword_alignment"),
    }

    for label, value in checklist_items.items():
        st.checkbox(label, value=bool(value), disabled=True)

    st.download_button(
        "⬇️ Download Analysis as JSON",
        data=json.dumps(result, indent=2),
        file_name="resume_ats_analysis.json",
        mime="application/json",
    )

st.divider()
st.caption(
    "ATS score is an AI-generated heuristic for resume improvement and should "
    "not be treated as an official score from a specific Applicant Tracking System."
)
